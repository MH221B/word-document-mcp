"""Build the sample documents with python-docx.

plain.docx needs: headings, a repeated word to make matching ambiguous, a list,
and an image so the media tools have something to act on.
"""

import base64

from docx import Document
from docx.shared import Inches

# a valid 1x1 PNG, so no binary asset needs committing
DOT_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAAC0lEQVR42mNkYAAAAAYAAjCB0C8AAAAASUVORK5CYII="
)


def plain() -> None:
    with open("samples/dot.png", "wb") as handle:
        handle.write(DOT_PNG)
    doc = Document()
    doc.add_heading("Overview", level=1)
    doc.add_paragraph("Total riders: 8")
    doc.add_picture("samples/dot.png", width=Inches(1.5))
    doc.add_heading("Risks", level=1)
    doc.add_paragraph("Total risks: 3")
    doc.add_paragraph("First item", style="List Bullet")
    doc.add_paragraph("Second item", style="List Bullet")
    doc.add_paragraph("Totals")
    doc.save("samples/plain.docx")


def template() -> None:
    doc = Document()
    doc.add_heading("Stage report", level=1)
    doc.add_paragraph("Summary text on the template style.")
    para = doc.add_paragraph("Hand edited paragraph")
    para.runs[0].bold = True
    doc.save("samples/template.docx")


if __name__ == "__main__":
    plain()
    template()
